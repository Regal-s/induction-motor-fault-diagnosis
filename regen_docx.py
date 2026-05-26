r"""
Regenerate manuscript_ieee_final.docx from the (rewritten) LaTeX source so the Word
version matches main.tex. Pipeline:
  1. Rasterize manuscript/figs/*.pdf -> manuscript/figs_png/*.png (Word can't embed PDF).
  2. Make a temp .tex whose \includegraphics point at the PNGs.
  3. pandoc temp.tex -> manuscript_ieee_final.docx (equations -> OMML, tables, bibliography).
  4. Post-process with python-docx: two-column layout + cap image widths to the column.
"""
import re, shutil, subprocess
from pathlib import Path
import fitz  # PyMuPDF
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# pandoc cannot parse the grouped-header \multicolumn table*, so we re-insert it.
BENCH_CAP = ("Full model benchmark on identical leakage-safe splits: StratifiedGroupKFold (GK, "
             "in-distribution) and Leave-One-Load-Out (LOLO, cross-load). Detection and phase report "
             "macro-F1; severity reports within-one-level accuracy. Best per column in bold. "
             "TabPFN-2.5 and xLSTM are added in this work.")
BENCH_ROWS = [
    ("Logistic Regression", [0.978, 0.978, 0.997, 0.971, 0.970, 0.920]),
    ("SVM-RBF",             [0.978, 0.977, 0.991, 0.975, 0.984, 0.967]),
    ("k-NN",                [0.977, 0.977, 0.937, 0.942, 0.982, 0.981]),
    ("MLP",                 [0.978, 0.978, 0.997, 0.982, 0.956, 0.981]),
    ("Random Forest",       [0.978, 0.979, 0.985, 0.688, 0.965, 0.997]),
    ("XGBoost",             [0.976, 0.948, 0.990, 0.943, 0.976, 0.997]),
    ("TabPFN-2.5",          [0.979, 0.979, 0.997, 0.969, 0.968, 0.991]),
    ("xLSTM",               [0.957, 0.890, 0.982, 0.908, 0.921, 0.961]),
]
BENCH_MAX = [max(r[1][j] for r in BENCH_ROWS) for j in range(6)]


def _cell(cell, text, bold=False, center=True):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text); run.font.size = Pt(8); run.font.bold = bold


def insert_benchmark_table(d):
    anchor = None
    for p in d.paragraphs:
        t = p.text
        if "benchmark eight models" in t or "in-distribution task is saturated" in t:
            anchor = p; break
    if anchor is None:        # fall back to the figure caption
        for p in d.paragraphs:
            if "Model benchmark across the three diagnostic tasks" in p.text:
                anchor = p; break
    if anchor is None:
        print("WARN: benchmark anchor not found; table not inserted"); return
    cur = anchor._p
    cap = d.add_paragraph("TABLE: " + BENCH_CAP)
    for r in cap.runs:
        r.font.size = Pt(8); r.font.bold = True
    cur.addnext(cap._p); cur = cap._p

    tbl = d.add_table(rows=2 + len(BENCH_ROWS), cols=7)
    try:
        tbl.style = "Table Grid"
    except Exception:
        pass
    r0 = tbl.rows[0].cells; r1 = tbl.rows[1].cells
    _cell(r0[0], "")
    _cell(r0[1].merge(r0[2]), "Detection (macro-F1)", bold=True)
    _cell(r0[3].merge(r0[4]), "Severity (within-1)", bold=True)
    _cell(r0[5].merge(r0[6]), "Phase (macro-F1)", bold=True)
    for i, lab in enumerate(["Model", "GK", "LOLO", "GK", "LOLO", "GK", "LOLO"]):
        _cell(r1[i], lab, bold=True, center=(i != 0))
    for ri, (name, vals) in enumerate(BENCH_ROWS, start=2):
        cells = tbl.rows[ri].cells
        _cell(cells[0], name, center=False)
        for j, v in enumerate(vals):
            _cell(cells[j + 1], f"{v:.3f}", bold=(abs(v - BENCH_MAX[j]) < 1e-9))
    cur.addnext(tbl._tbl)
    print("inserted 8-model benchmark table")

ROOT = Path(r"D:\Naveen")
MAN = ROOT / "manuscript"
FIGS = MAN / "figs"
PNG = MAN / "figs_png"
OUT = ROOT / "manuscript_ieee_final.docx"
DPI = 200

def rasterize():
    PNG.mkdir(exist_ok=True)
    for pdf in FIGS.glob("*.pdf"):
        d = fitz.open(pdf)
        d[0].get_pixmap(dpi=DPI).save(PNG / (pdf.stem + ".png"))
        d.close()
    for p in FIGS.glob("*.png"):           # carry over already-raster figures
        shutil.copyfile(p, PNG / p.name)
    print(f"rasterized -> {PNG} ({len(list(PNG.glob('*.png')))} pngs)")

def make_tex():
    tex = (MAN / "main.tex").read_text(encoding="utf-8")
    tex = re.sub(r"figs/([A-Za-z0-9_]+)\.(pdf|png)", r"figs_png/\1.png", tex)
    tmp = MAN / "main_fordocx.tex"
    tmp.write_text(tex, encoding="utf-8")
    return tmp

def run_pandoc(tmp):
    if OUT.exists():
        shutil.copyfile(OUT, OUT.with_suffix(".prerewrite.bak.docx"))  # backup
    cmd = ["pandoc", tmp.name, "-o", str(OUT), "--resource-path", str(MAN),
           "--mathml", "-f", "latex+raw_tex"]
    r = subprocess.run(cmd, cwd=str(MAN), capture_output=True, text=True)
    print("pandoc rc=", r.returncode)
    if r.stderr.strip():
        print(r.stderr.strip()[:1500])
    return r.returncode == 0

def two_columns_and_fit():
    d = docx.Document(str(OUT))
    col_w = Inches(3.2)
    for shp in d.inline_shapes:
        if shp.width and shp.width > col_w:
            ratio = int(col_w) / shp.width
            shp.height = int(shp.height * ratio)
            shp.width = int(col_w)
    insert_benchmark_table(d)
    sectPr = d.sections[0]._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); sectPr.append(cols)
    cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "432")
    d.save(str(OUT))
    d2 = docx.Document(str(OUT))
    print(f"saved {OUT.name}: paragraphs={len(d2.paragraphs)} tables={len(d2.tables)} "
          f"images={len(d2.inline_shapes)}")

if __name__ == "__main__":
    rasterize()
    tmp = make_tex()
    if run_pandoc(tmp):
        two_columns_and_fit()
    tmp.unlink(missing_ok=True)
